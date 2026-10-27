"""Build the Café Poliçe article (Turkish) as .docx and .md from the engine's outputs."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
M = json.loads((ROOT / "outputs" / "metrics.json").read_text(encoding="utf-8"))
E = json.loads((ROOT / "outputs" / "evt_uncertainty.json").read_text(encoding="utf-8"))
GITHUB = "https://github.com/000MAK/cafepolice-actuarial-loss-engine"
HF = "https://huggingface.co/makactuary/fremtpl2-pricing-engine"
NAVY, ORANGE = RGBColor(0x14, 0x21, 0x3D), RGBColor(0xB4, 0x56, 0x1C)


def tr(x: float, d: int = 0) -> str:
    """Turkish number format: thousands '.', decimals ','."""
    s = f"{x:,.{d}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


c, f, s, pp = M["cleaning"], M["frequency"], M["severity"], M["pure_premium"]
gpd = s["gpd"]

TITLE = "Bir Yılın Faturasını On Hasar Yazar"
SUBTITLE = "Açık veriyle bir trafik sigortası fiyatlama motoru: frekans, şiddet ve kuyruk riski"
BYLINE = "Mehmet Ali Kaya · Veri Yönetimi Uzmanı, Yönetim Bilişim Sistemleri Doktora Adayı, İstatistik Yüksek Lisans"
SPOT = (f"{tr(c['policies'])} poliçelik açık bir trafik portföyünde hasar frekansını, hasar şiddetini ve büyük hasarları ayrı ayrı modelledim. "
        "Makine öğrenmesi riski daha iyi sıraladı; ama asıl ders, yılın faturasının birkaç büyük hasar tarafından yazılmasıydı.")

SECTIONS = [
    ("Hasarın iki sorusu", [
        "Her sigortacı bir poliçeye bakarken aslında iki ayrı soru sorar: Bu poliçe hasar yapar mı, yaparsa ne kadar tutar? "
        "Aktüerya dilinde bunlara frekans ve şiddet diyoruz; saf prim de bu ikisinin çarpımından doğuyor. Kâğıt üzerinde sade bir çarpma işlemi. "
        "Sahada ise bu iki sorunun karakteri birbirine hiç benzemiyor: Frekans, kalabalığın davranışını anlatan sakin bir hikâye; "
        "şiddet ise ara sıra sahneye çıkıp bütün yılın sonucunu değiştiren birkaç büyük hasarın hikâyesi.",
        "Bu yazıda, herkesin indirip tekrarlayabileceği açık bir veri setiyle bu iki hikâyeyi ayrı ayrı modelledim, "
        "sonra yeniden birleştirip bir fiyatlama motoru kurdum. Kodun tamamı, sonuçlar ve modelin kendisi açık; yazının sonunda bağlantıları bulabilirsiniz.",
    ]),
    ("Veri: 678 bin poliçe, 26 bin hasar", [
        f"Kullandığım veri, aktüerya eğitiminde sıkça başvurulan freMTPL2: Fransız bir sigortacının trafik (motor third-party liability) portföyü. "
        f"Veri, Christophe Dutang ve Arthur Charpentier'in CASdatasets koleksiyonunun parçası ve OpenML üzerinde CC0 lisansıyla, yani kısıtsız olarak yayımlanıyor. "
        f"Temizlik sonrası elimde {tr(c['policies'])} poliçe, toplam {tr(c['total_exposure'])} poliçe-yılı maruziyet ve {tr(c['claims'])} hasar vardı. "
        f"Portföyün yıllık hasar frekansı yüzde {tr(c['portfolio_frequency']*100, 1)}; yani yaklaşık her 14 poliçe-yılında bir hasar.",
        f"Her açık veri setinin bir hikâyesi olduğu gibi kusurları da var ve bunları saklamamak gerekir. Maruziyeti bir yılı aşan {tr(c['exposure_capped'])} kaydı bir yılla sınırladım, "
        f"poliçe başına hasar sayısını 4'te kestim. Daha önemlisi, {tr(c['claims_without_amount'])} poliçede hasar sayısı görünüyor ama hasar tutarı kaydı yok. "
        "Bu tutarsızlık veri setinin bilinen bir özelliği; scikit-learn dokümantasyonundaki yaklaşımla bu poliçeleri hasarsız kabul ettim. "
        "Bu tür kararlar sonucu etkiler; o yüzden her birini koda ve yazıya açıkça yazdım.",
    ]),
    ("Frekans: Kalabalığın davranışı", [
        "Önce verinin kendisine bakalım. Bonus-malus katsayısı 100'ün üzerindeki sürücülerin yıllık hasar frekansı, 50'deki (en iyi basamak) sürücülerin yaklaşık 6,9 katı. "
        "18-21 yaş aralığındaki sürücüler ise 40-50 yaş grubuna göre 2,8 kat daha sık hasar yapıyor (Grafik 1). Bu, sektörün sezgisiyle örtüşen ama rakamla görmesi hep öğretici bir tablo.",
        f"Frekansı iki yöntemle modelledim: Sektörün iş atı olan Poisson GLM ve bir makine öğrenmesi yöntemi olan Poisson kayıplı gradient boosting. "
        f"İkisini de verinin yüzde 80'iyle eğittim, hiç görmedikleri yüzde 20'lik bir test portföyünde ({tr(M['test_policies'])} poliçe) ölçtüm. "
        f"Riski sıralama gücünü gösteren Gini katsayısı GLM'de {tr(f['poisson_glm']['gini'], 3)}, gradient boosting'de {tr(f['poisson_gbm']['gini'], 3)} çıktı (Grafik 2). "
        f"Açıklanan sapma (D²) ise sırasıyla yüzde {tr(f['poisson_glm']['d2_vs_mean']*100, 1)} ve yüzde {tr(f['poisson_gbm']['d2_vs_mean']*100, 1)}.",
        "Bu D² değerleri ilk bakışta düşük görünebilir; ama trafik hasarında beklenen bu. Tek bir sürücünün o yıl kaza yapıp yapmayacağı büyük ölçüde şansa bağlı; "
        "modelin işi bireyi tahmin etmek değil, riskleri doğru sıraya dizmek. Gradient boosting bu sıralamayı belirgin biçimde daha iyi yapıyor. "
        "Bedeli ise açıklanabilirlik: GLM'in katsayıları tek tek okunabilirken, boosting modelini anlatmak için ek araçlar gerekiyor. "
        "Düzenleyicinin ve müşterinin sorduğu \"neden bu fiyat?\" sorusu, sadece doğruluk kadar önemli.",
    ]),
    ("Şiddet: Uzun kuyruğun ağırlığı", [
        f"Hasar tutarlarına geçince tablo değişiyor. Hasar başına tutarın medyanı 1.172 avro, ortalaması 2.279 avro; ortalamanın medyanın neredeyse iki katı olması, "
        f"dağılımın sağa doğru upuzun bir kuyruğu olduğunu söylüyor. Veri setindeki en büyük tek hasar 4,07 milyon avro (Grafik 3).",
        f"Bu yüzden şiddeti tek bir modelle değil, iki parçada ele aldım. Dağılımın yüzde 99'luk dilimine ({tr(s['threshold_u_p99'])} avro) bir eşik koydum. "
        f"Eşiğin altındaki olağan hasarları Gamma GLM ile, eşiği aşan kısmı ise Uç Değer Teorisi'nin (EVT) Genelleştirilmiş Pareto dağılımıyla modelledim. "
        f"Eşiği aşan {s['claims_above_u']} hasara uyan dağılımın şekil parametresi ξ = {tr(gpd['xi'], 2)} çıktı.",
        f"Bu sayı yazının belki de en önemli bulgusu. ξ değeri 1'e yaklaştıkça kuyruk ağırlaşır; 1'in üzerinde ise dağılımın ortalaması bile tanımsız hale gelir. "
        f"Bootstrap yöntemiyle ξ için hesapladığım yüzde 95 güven aralığı {tr(E['xi_ci95'][0], 2)} ile {tr(E['xi_ci95'][1], 2)} arasında; "
        f"yeniden örneklemelerin yüzde {tr(E['share_xi_ge_1']*100, 0)}'sinde ξ 1'i aştı. Eşiği değiştirdiğimde de ξ 0,74 ile 1,00 arasında gezindi (Grafik 4). "
        "Kısacası bu portföyde büyük hasarın \"ortalama maliyeti\" diye sağlam bir sayı yok; yalnızca bir aralık var.",
        f"Bu belirsizliğin pratikteki karşılığını test portföyünde gördüm: En büyük 10 hasar, test yılının toplam hasar maliyetinin yüzde {tr(E['test_top10_share']*100, 0)}'sini, "
        "en büyük 100 hasar ise yüzde 41'ini oluşturuyor (Grafik 5). Bir yılın faturasını, gerçekten de bir avuç hasar yazıyor.",
    ]),
    ("Motoru birleştirmek: Saf prim", [
        f"Fiyatlama motoru üç parçadan oluşuyor: Gradient boosting ile frekans, Gamma GLM ile olağan hasar şiddeti ve her hasara eklenen bir büyük hasar yüklemesi. "
        f"İşin can alıcı noktası bu son kalem. Gözlenen büyük hasarlardan hesaplanan yükleme hasar başına {tr(E['empirical_load'])} avro; "
        f"GPD dağılımının ima ettiği yükleme ise {tr(s['large_loss_load_per_claim'])} avro. Bootstrap bu ikincisi için 509 ile 4.917 avro arasında bir aralık veriyor.",
        f"Test portföyünde gözleme dayalı yüklemeyle kurulan motor, toplam hasarı neredeyse kusursuz tahmin etti: tahmin/gerçekleşen oranı {tr(E['pp_ratio_empirical_load'], 3)}. "
        f"GPD yüklemesiyle bu oran {tr(E['pp_ratio_gpd_load'], 3)}'e çıkıyor; yani prim yüzde 23 daha ihtiyatlı. "
        f"Karşılaştırma için tek adımda kurulan bir Tweedie GLM'in oranı {tr(pp['tweedie_glm']['predicted_over_observed'], 2)} ve Gini'si {tr(pp['tweedie_glm']['gini'], 3)}; "
        f"frekans ve şiddeti ayrı modelleyen yaklaşımın Gini'si ise {tr(pp['freq_gbm_x_sev']['gini'], 3)}.",
        "Risk dilimlerine bakınca (Grafik 6) motorun en düşük ve en yüksek risk dilimlerinde gerçekleşenin yüzde 15-20 kadar altında tahmin yaptığı görülüyor; "
        "ara dilimlerde ise sapma iki yönde de var. Toplamın dengede olması, her dilimin dengede olduğu anlamına gelmiyor; bu, fiyatlamada ayrıca izlenmesi gereken bir nokta.",
    ]),
    ("Sektör için üç ders", [
        "Birincisi, makine öğrenmesi riski daha iyi sıralıyor ama tek başına bir tarife değil. Gradient boosting'in kazandırdığı Gini puanları gerçek; "
        "ancak bu modeli tarifeye taşımak için açıklanabilirlik, kararlılık ve düzenleyici kabul gibi soruların da cevaplanması gerekiyor. "
        "GLM ile makine öğrenmesini rakip değil, birbirini denetleyen iki göz olarak düşünmek bence en sağlıklısı.",
        "İkincisi, büyük hasar yüklemesi bir matematik sorusu değil, bir risk iştahı kararı. Kuyruk bu kadar ağır olduğunda geçmiş deneyim gelecekteki büyük hasarı "
        "sistematik olarak eksik gösterebilir; EVT ise ihtiyatlı ama belirsiz bir rakam verir. Hangisinin seçileceği; şirketin sermayesine, reasürans yapısına ve fiyat rekabetine bağlı. "
        "Model bu kararı vermez, kararın bedelini görünür kılar.",
        "Üçüncüsü, veri kalitesi her şeyin önünde geliyor. Bu veri setinde 9 binden fazla poliçede hasar sayısı ile hasar tutarı birbirini tutmuyordu; "
        "bu kayıtları nasıl ele aldığınız, saf primi doğrudan değiştiriyor. Veri yönetimi, fiyatlamanın arka odası değil, temelidir.",
    ]),
    ("Sınırlılıklar", [
        "Bu çalışma Fransız bir portföyün verisine dayanıyor; Türkiye trafik sigortasındaki tarife yapısı, enflasyon, kur etkisi ve hasar dinamikleri farklı. "
        "Sonuçlar bir yöntem gösterimi olarak okunmalı, doğrudan bir tarife önerisi olarak değil. Veri tek bir döneme ait; hasar enflasyonu, gecikmeli bildirilen hasarlar (IBNR) "
        "ve reasürans etkisi modelde yer almıyor. Büyük hasar modeli 199 hasara dayanıyor; bu kadar az gözlemle kuyruk hakkında söylenebilecekler sınırlı.",
    ]),
]

FIGURES = {  # inserted after the named section
    "Frekans: Kalabalığın davranışı": [("01_frekans.png", "Grafik 1. Bonus-malus ve sürücü yaşına göre gözlenen yıllık hasar frekansı."),
                                       ("04_lorenz.png", "Grafik 2. Test portföyünde Lorenz eğrileri: eğri köşegenden ne kadar uzaksa model riski o kadar iyi sıralıyor.")],
    "Şiddet: Uzun kuyruğun ağırlığı": [("02_siddet.png", "Grafik 3. Hasar başına tutarın dağılımı (logaritmik ölçek) ve yüzde 99 büyük hasar eşiği."),
                                       ("03_kuyruk.png", "Grafik 4. Eşiğe göre GPD şekil parametresi ξ; turuncu çubuk yüzde 99 eşiğindeki yüzde 95 bootstrap güven aralığı."),
                                       ("06_yogunlasma.png", "Grafik 5. Test yılında hasar maliyetinin en büyük hasarlarda yoğunlaşması.")],
    "Motoru birleştirmek: Saf prim": [("05_kalibrasyon.png", "Grafik 6. Saf prim kalibrasyonu: 10 risk diliminde tahmin edilen ve gerçekleşen hasar maliyeti.")],
}


def bottom_border(paragraph, color="E07A2E") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.2
    for level, size in ((1, 15),):
        h = doc.styles[f"Heading {level}"]
        h.font.name, h.font.size, h.font.color.rgb, h.font.bold = "Calibri", Pt(size), NAVY, True
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(14), Pt(6)

    t = doc.add_paragraph()
    r = t.add_run(TITLE)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(26), NAVY
    st = doc.add_paragraph()
    r = st.add_run(SUBTITLE)
    r.font.size, r.font.color.rgb = Pt(13), ORANGE
    by = doc.add_paragraph()
    r = by.add_run(BYLINE)
    r.italic, r.font.size = True, Pt(9.5)
    bottom_border(by)
    sp = doc.add_paragraph()
    r = sp.add_run(SPOT)
    r.bold, r.font.size = True, Pt(11.5)

    for heading, paras in SECTIONS:
        doc.add_heading(heading, level=1)
        for text in paras:
            doc.add_paragraph(text)
        if heading == "Şiddet: Uzun kuyruğun ağırlığı":
            q = doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(f"“En büyük 10 hasar, test yılının toplam maliyetinin yüzde {tr(E['test_top10_share']*100, 0)}'sini oluşturdu.”")
            r.italic, r.bold, r.font.size, r.font.color.rgb = True, True, Pt(14), ORANGE
        for img, caption in FIGURES.get(heading, []):
            doc.add_picture(str(ROOT / "outputs" / "charts" / img), width=Cm(9.5 if img.startswith("04_") else 15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(caption)
            r.italic, r.font.size = True, Pt(9)
        if heading == "Motoru birleştirmek: Saf prim":
            rows = [("Model", "Gini", "Tahmin / gerçekleşen"),
                    ("Frekans (GBM) × şiddet, gözlenen büyük hasar yüklemesi", tr(pp["freq_gbm_x_sev"]["gini"], 3), tr(E["pp_ratio_empirical_load"], 3)),
                    ("Frekans (GBM) × şiddet, GPD büyük hasar yüklemesi", tr(pp["freq_gbm_x_sev"]["gini"], 3), tr(E["pp_ratio_gpd_load"], 3)),
                    ("Frekans (GLM) × şiddet, GPD yüklemesi", tr(pp["freq_glm_x_sev"]["gini"], 3), tr(pp["freq_glm_x_sev"]["predicted_over_observed"], 3)),
                    ("Tek adımlı Tweedie GLM", tr(pp["tweedie_glm"]["gini"], 3), tr(pp["tweedie_glm"]["predicted_over_observed"], 3))]
            table = doc.add_table(rows=len(rows), cols=3)
            table.style = "Light Grid Accent 1"
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cell = table.cell(i, j)
                    cell.text = val
                    cell.paragraphs[0].runs[0].font.size = Pt(9.5)
                    cell.paragraphs[0].runs[0].bold = i == 0
            for row in table.rows:
                row.cells[0].width, row.cells[1].width, row.cells[2].width = Cm(10), Cm(2.6), Cm(3.6)
            cp = doc.add_paragraph()
            r = cp.add_run(f"Tablo 1. Test portföyünde ({tr(M['test_policies'])} poliçe) saf prim modelleri.")
            r.italic, r.font.size = True, Pt(9)

    doc.add_heading("Kod, veri ve model", level=1)
    doc.add_paragraph("Bu yazıdaki her sayı, aşağıdaki açık kaynak koddan yeniden üretilebilir:")
    for label, url in (("Kod ve sonuçlar (GitHub): ", GITHUB), ("Model ve metrikler (Hugging Face): ", HF)):
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(url)
    src = doc.add_paragraph()
    r = src.add_run("Kaynak: Dutang, C. ve Charpentier, A., CASdatasets: Insurance datasets (freMTPL2freq ve freMTPL2sev, OpenML 41214 ve 41215, CC0). "
                    "Yöntemsel çerçeve için bkz. Charpentier, A. (ed.), Computational Actuarial Science with R, CRC Press.")
    r.italic, r.font.size = True, Pt(9)
    doc.save(path)


def build_md(path: Path) -> None:
    lines = [f"# {TITLE}", f"*{SUBTITLE}*", "", f"*{BYLINE}*", "", f"**{SPOT}**", ""]
    for heading, paras in SECTIONS:
        lines += [f"## {heading}", ""] + [p + "\n" for p in paras]
        for img, caption in FIGURES.get(heading, []):
            lines += [f"![{caption}](../outputs/charts/{img})", f"*{caption}*", ""]
    lines += ["## Kod, veri ve model", "", f"- Kod ve sonuçlar: {GITHUB}", f"- Model: {HF}", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    out = ROOT / "article"
    build_docx(out / "CafePolice_HasarAnalizi_MehmetAliKaya.docx")
    build_md(out / "CafePolice_HasarAnalizi_MehmetAliKaya.md")
    print("written:", sorted(p.name for p in out.iterdir()))
